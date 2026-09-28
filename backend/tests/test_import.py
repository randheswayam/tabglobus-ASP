"""Bulk import of in-progress projects (PRD 7.19, FR-27): a preview validates every row and writes nothing."""

import pytest

from app.models import Project

HEADER = (
    "project_name,location,client_name,client_email,site_address,current_stage,"
    "civil_engineer_email,confirmed_by,legal_expected_date"
)


def _csv(*rows: str) -> bytes:
    return ("\n".join([HEADER, *rows]) + "\n").encode()


def _preview(client, headers, data: bytes, name="projects.csv", content_type="text/csv"):
    return client.post("/admin/import/projects/preview", files={"file": (name, data, content_type)}, headers=headers)


GOOD = (
    "Sathe House,Pashan Pune,Sathe family,sathe@example.com,Survey 12 Pashan,grid,"
    "engineer@siteflow.local,Parvez,2026-12-01"
)
NEW = "Mane Villa,Wakad Pune,,,,,engineer@siteflow.local,,"


def test_template_lists_the_columns(client, auth_headers):
    r = client.get("/admin/import/projects/template", headers=auth_headers("admin"))
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.text.splitlines()[0] == HEADER


def test_preview_reports_good_rows_and_writes_nothing(client, auth_headers, db):
    r = _preview(client, auth_headers("admin"), _csv(GOOD, NEW))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["counts"] == {"rows": 2, "ok": 2, "errors": 0}
    first = body["rows"][0]
    assert first["row"] == 2 and first["ok"] is True and first["errors"] == []
    assert first["values"]["project_name"] == "Sathe House" and first["values"]["current_stage"] == "grid"
    assert body["rows"][1]["values"]["current_stage"] == ""  # blank: a new project from Project setup
    assert db.query(Project).count() == 0


@pytest.mark.parametrize(
    "row,error",
    [
        ("Villa,Pune,,,,roofing,engineer@siteflow.local,Parvez,", "current_stage 'roofing' is not a stage"),
        ("Villa,Pune,,,,grid,nobody@example.com,Parvez,", "civil_engineer_email is not an active Civil Engineer"),
        ("Villa,Pune,,,,grid,architect@siteflow.local,Parvez,", "civil_engineer_email is not an active Civil Engineer"),
        (",Pune,,,,,engineer@siteflow.local,,", "project_name is required"),
        ("Villa,,,,,,engineer@siteflow.local,,", "location is required"),
        ("Villa,Pune,,,,grid,engineer@siteflow.local,,", "confirmed_by is required when current_stage is set"),
        ("Villa,Pune,,,,setup,engineer@siteflow.local,Parvez,", "current_stage 'setup' is the first stage"),
        ("Villa,Pune,,bad-email,,,engineer@siteflow.local,,", "client_email is not an email address"),
        ("Villa,Pune,,,,,engineer@siteflow.local,,31-12-2026", "legal_expected_date must be YYYY-MM-DD"),
        ("=HYPERLINK(1),Pune,,,,,engineer@siteflow.local,,", "project_name starts with = + - or @"),
        ("Villa,@Pune,,,,,engineer@siteflow.local,,", "location starts with = + - or @"),
    ],
)
def test_row_errors(client, auth_headers, row, error):
    body = _preview(client, auth_headers("admin"), _csv(row)).json()
    assert body["counts"]["errors"] == 1
    assert any(error in e for e in body["rows"][0]["errors"]), body["rows"][0]["errors"]


def test_duplicate_names_in_the_file_and_in_siteflow(client, auth_headers, new_project):
    new_project("Existing Villa")
    body = _preview(
        client,
        auth_headers("admin"),
        _csv(NEW, NEW.replace("Mane Villa", "mane villa"), NEW.replace("Mane Villa", "Existing Villa")),
    ).json()
    assert body["rows"][0]["ok"] is True
    assert "project_name repeats row 2" in body["rows"][1]["errors"]
    assert "a project named 'Existing Villa' already exists" in body["rows"][2]["errors"]


def test_file_level_checks(client, auth_headers):
    admin = auth_headers("admin")
    assert _preview(client, admin, b"name,place\nVilla,Pune\n").status_code == 422  # wrong header
    assert _preview(client, admin, _csv(*[NEW.replace("Mane", f"M{i}") for i in range(501)])).status_code == 422
    assert _preview(client, admin, b"x" * (1024 * 1024 + 1)).status_code == 413
    assert _preview(client, admin, b"\x89PNG\r\n\x1a\n", name="p.png", content_type="image/png").status_code == 415
    assert _preview(client, admin, b"\xff\xfe\x00bad").status_code == 422  # not UTF-8 text
    assert _preview(client, admin, HEADER.encode() + b"\n").json()["counts"]["rows"] == 0


def test_a_utf8_bom_from_excel_is_accepted(client, auth_headers):
    r = _preview(client, auth_headers("admin"), b"\xef\xbb\xbf" + _csv(NEW))
    assert r.status_code == 200 and r.json()["counts"]["ok"] == 1


@pytest.mark.parametrize("role", ["architect", "team_lead", "accounts"])
def test_only_admins_import(client, auth_headers, role):
    assert _preview(client, auth_headers(role), _csv(NEW)).status_code == 403
    assert client.get("/admin/import/projects/template", headers=auth_headers(role)).status_code == 403


# ---------- commit (Task 21) ----------

BAD = "Villa Bad,Pune,,,,roofing,engineer@siteflow.local,Parvez,"


def _commit(client, headers, data: bytes, mode: str, architect_id, name="projects.csv"):
    return client.post(
        "/admin/import/projects/commit",
        files={"file": (name, data, "text/csv")},
        data={"mode": mode, "architect_id": str(architect_id)},
        headers=headers,
    )


def test_all_or_nothing_imports_nothing_when_a_row_is_bad(client, auth_headers, users, db):
    from app.models import ImportBatch

    r = _commit(client, auth_headers("admin"), _csv(GOOD, BAD), "all_or_nothing", users["architect"].id)
    assert r.status_code == 422
    assert r.json()["detail"]["counts"] == {"rows": 2, "ok": 1, "errors": 1}
    assert db.query(Project).count() == 0
    (batch,) = db.query(ImportBatch).all()
    assert batch.mode == "all_or_nothing" and batch.imported == 0 and batch.rows == 2
    assert batch.errors == [{"row": 3, "errors": ["current_stage 'roofing' is not a stage"]}]


def test_valid_rows_only_imports_the_good_rows_through_normal_onboarding(client, auth_headers, users, db):
    from app.models import AuditEvent, ImportBatch

    admin = auth_headers("admin")
    r = _commit(client, admin, _csv(GOOD, BAD, NEW), "valid_rows_only", users["architect"].id)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["counts"] == {"rows": 3, "ok": 2, "errors": 1}
    assert [p["name"] for p in body["projects"]] == ["Sathe House", "Mane Villa"]
    batch = db.query(ImportBatch).one()
    assert batch.imported == 2 and len(batch.sha256) == 64 and batch.filename == "projects.csv"
    assert batch.by_id == users["admin"].id

    arch = auth_headers("architect")
    sathe = client.get(f"/projects/{body['projects'][0]['id']}", headers=arch).json()
    assert sathe["client"]["name"] == "Sathe family" and sathe["site"]["address"] == "Survey 12 Pashan"
    assert sathe["legal_approval"]["expected_date"] == "2026-12-01"
    view = client.get(f"/projects/{sathe['id']}/stages", headers=arch).json()
    states = {s["key"]: s for p in view["phases"] for s in p["stages"]}
    assert states["investigations"]["state"] == "historical"
    assert states["investigations"]["historical"]["confirmed_by"] == "Parvez"
    assert states["grid"]["state"] == "active"
    mane = client.get(f"/projects/{body['projects'][1]['id']}/stages", headers=arch).json()
    assert mane["current_stages"] == ["Project setup"]
    created = db.query(AuditEvent).filter_by(action="project.created").all()
    assert {e.detail["import_batch_id"] for e in created} == {batch.id}
    assert db.query(AuditEvent).filter_by(action="import.committed").one().detail["imported"] == 2


def test_all_or_nothing_imports_everything_when_every_row_is_good(client, auth_headers, users, db):
    r = _commit(client, auth_headers("admin"), _csv(GOOD, NEW), "all_or_nothing", users["architect"].id)
    assert r.status_code == 201 and r.json()["counts"]["ok"] == 2
    assert db.query(Project).count() == 2


def test_the_architect_must_be_an_active_architect(client, auth_headers, users):
    admin = auth_headers("admin")
    assert _commit(client, admin, _csv(NEW), "valid_rows_only", users["team_lead"].id).status_code == 422
    assert _commit(client, admin, _csv(NEW), "valid_rows_only", 9999).status_code == 422
    assert _commit(client, admin, _csv(NEW), "sometimes", users["architect"].id).status_code == 422


def test_batches_are_listed_for_admins(client, auth_headers, users):
    admin = auth_headers("admin")
    _commit(client, admin, _csv(NEW), "valid_rows_only", users["architect"].id)
    r = client.get("/admin/import/batches", headers=admin)
    assert r.status_code == 200
    (b,) = r.json()
    assert b["filename"] == "projects.csv" and b["imported"] == 1 and b["by"]["name"] == users["admin"].name
    assert client.get("/admin/import/batches", headers=auth_headers("architect")).status_code == 403


# ---------- XLSX (sprint v4 Task 41): the same cases, read from the first sheet ----------


def _xlsx(*rows: str, sheet_first=None) -> bytes:
    import csv
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if sheet_first is not None:  # a decoy sheet placed after the real one must be ignored
        wb.create_sheet("Notes").append(sheet_first)
    for line in csv.reader(io.StringIO("\n".join([HEADER, *rows]))):
        ws.append([v if v != "" else None for v in line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_an_xlsx_previews_like_the_csv(client, auth_headers):
    admin = auth_headers("admin")
    x = _preview(client, admin, _xlsx(GOOD, NEW, sheet_first=["ignore me"]), name="projects.xlsx", content_type=XLSX)
    c = _preview(client, admin, _csv(GOOD, NEW))
    assert x.status_code == 200, x.text
    assert x.json() == c.json()


@pytest.mark.parametrize(
    "row,error",
    [
        ("Villa,Pune,,,,roofing,engineer@siteflow.local,Parvez,", "current_stage 'roofing' is not a stage"),
        (",Pune,,,,,engineer@siteflow.local,,", "project_name is required"),
        ("Villa,Pune,,,,,engineer@siteflow.local,,31-12-2026", "legal_expected_date must be YYYY-MM-DD"),
        ("=HYPERLINK(1),Pune,,,,,engineer@siteflow.local,,", "project_name starts with = + - or @"),
    ],
)
def test_xlsx_row_errors(client, auth_headers, row, error):
    body = _preview(client, auth_headers("admin"), _xlsx(row), name="p.xlsx", content_type=XLSX).json()
    assert any(error in e for e in body["rows"][0]["errors"]), body


def test_an_excel_date_cell_reads_as_iso(client, auth_headers):
    import datetime
    import io

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(HEADER.split(","))
    wb.active.append(
        [
            "Kale Villa",
            "Baner Pune",
            None,
            None,
            None,
            None,
            "engineer@siteflow.local",
            None,
            datetime.datetime(2026, 12, 1),
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    body = _preview(client, auth_headers("admin"), buf.getvalue(), name="d.xlsx", content_type=XLSX).json()
    assert body["rows"][0]["ok"] is True and body["rows"][0]["values"]["legal_expected_date"] == "2026-12-01"


def test_an_xlsx_commits_through_normal_onboarding(client, auth_headers, users, db):
    from app.models import ImportBatch

    r = _commit(
        client, auth_headers("admin"), _xlsx(GOOD, NEW), "all_or_nothing", users["architect"].id, name="projects.xlsx"
    )
    assert r.status_code == 201, r.text
    assert db.query(Project).count() == 2 and db.query(ImportBatch).one().filename == "projects.xlsx"


def test_the_committed_xlsx_fixture_previews(client, auth_headers):
    from pathlib import Path

    data = (Path(__file__).parent / "fixtures" / "import.xlsx").read_bytes()
    body = _preview(client, auth_headers("admin"), data, name="import.xlsx", content_type=XLSX).json()
    assert body["counts"] == {"rows": 3, "ok": 2, "errors": 1}  # the fixture keeps one bad row, like the CSV


def test_the_file_type_is_checked_by_signature(client, auth_headers):
    admin = auth_headers("admin")
    # a CSV renamed to .xlsx, and a zip that isn't a workbook
    assert _preview(client, admin, _csv(NEW), name="p.xlsx", content_type=XLSX).status_code == 415
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "not a workbook")
    assert _preview(client, admin, buf.getvalue(), name="p.xlsx", content_type=XLSX).status_code == 422
    # a workbook renamed to .csv is refused as a CSV
    assert _preview(client, admin, _xlsx(NEW), name="p.csv").status_code == 415
